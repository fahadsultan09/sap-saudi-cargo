
from pathlib import Path
from openpyxl import load_workbook

DATA_DIR = Path("data/sap_excel")

FILES_AND_COLUMNS = {
    "LFBK - Vendor Master (Bank Details).xlsx": [
        "Supplier",
        "Bank Country/Region",
        "Bank Key",
    ],
    "EKKN - PO account assignment.xlsx": [
        "Purchasing Document",
        "Item",
        "G/L Account",
        "Cost Center",
        "Controlling Area",
    ],
    "ESSR - Service Entry Sheet Header Data.xlsx": [
        "Entry Sheet",
        "Purchasing Document",
        "Item",
        "Package number",
    ],
}

for filename, wanted_columns in FILES_AND_COLUMNS.items():
    path = DATA_DIR / filename
    print(f"\n{'=' * 70}\n{filename}")

    if not path.exists():
        print(f"FILE NOT FOUND: {path}")
        continue

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Data"] if "Data" in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)

    headers = [str(h).strip() if h is not None else "" for h in next(rows)]
    indexes = {
        column: headers.index(column)
        for column in wanted_columns
        if column in headers
    }

    print("Columns found:", list(indexes))

    for row_number, row in enumerate(rows, start=2):
        if row_number > 7:
            break
        print({
            column: row[index]
            for column, index in indexes.items()
        })

    wb.close()
