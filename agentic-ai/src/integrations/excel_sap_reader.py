from __future__ import annotations

import hashlib
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_DATA_DIR = (
    Path(__file__).resolve().parents[2] / "sap_excel"
)

FILES = {
    "ekko": "EKKO-PO HEADER.xlsx",
    "ekpo": "EKPO- PO LINE ITEM.xlsx",
    "ekbe": "EKBE - PO HISTORY.xlsx",
    "lfa1": "LFA1 - SUPPLIER MASTER DATA.xlsx",
}


def _text(value: Any) -> str | None:
    if value is None:
        return None
    result = str(value).strip()
    if not result:
        return None
    if isinstance(value, float) and value.is_integer():
        result = str(int(value))
    return result


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        result = Decimal(str(value).replace(",", "").strip())
        return result if result.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value

class ExcelSAPReadOnlyReader:
    """Read-only PO and supplier lookups from local SAP Excel exports."""

    def __init__(self, data_dir: str | Path | None = None):
        self.data_dir = Path(data_dir or DEFAULT_DATA_DIR).resolve()

        missing_files = [
            filename
            for filename in FILES.values()
            if not (self.data_dir / filename).is_file()
        ]
        if missing_files:
            raise FileNotFoundError(
                f"Missing Excel export(s) in {self.data_dir}: "
                + ", ".join(missing_files)
            )

    def _rows(
        self,
        key: str,
        required_columns: list[str],
    ) -> list[dict[str, Any]]:
        path = self.data_dir / FILES[key]

        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )
        try:
            worksheet = (
                workbook["Data"]
                if "Data" in workbook.sheetnames
                else workbook.active
            )
            rows = worksheet.iter_rows(values_only=True)
            headers = next(rows, None)

            if not headers:
                raise ValueError(f"No header row found in {path.name}")

            positions: dict[str, int] = {}
            for index, header in enumerate(headers):
                if header is not None:
                    positions.setdefault(str(header).strip(), index)

            missing = [
                column
                for column in required_columns
                if column not in positions
            ]
            if missing:
                raise ValueError(
                    f"{path.name} is missing required columns: {missing}"
                )

            result: list[dict[str, Any]] = []
            for values in rows:
                record = {
                    column: (
                        values[positions[column]]
                        if positions[column] < len(values)
                        else None
                    )
                    for column in required_columns
                }
                if any(value is not None for value in record.values()):
                    result.append(record)

            return result
        finally:
            workbook.close()

    def _snapshot_hash(self) -> str:
        digest = hashlib.sha256()

        for key in sorted(FILES):
            path = self.data_dir / FILES[key]
            digest.update(key.encode("utf-8"))
            digest.update(path.name.encode("utf-8"))

            with path.open("rb") as file:
                for chunk in iter(
                    lambda: file.read(1024 * 1024), b""
                ):
                    digest.update(chunk)

        return digest.hexdigest()

    def get_supplier(
        self,
        supplier_id: str,
    ) -> dict[str, Any] | None:
        requested_id = _text(supplier_id)
        if not requested_id:
            return None

        rows = self._rows(
            "lfa1",
            [
                "Supplier",
                "Name",
                "Name 2",
                "Country/Region Key",
                "City",
                "Central Deletion Flag",
                "Central posting block",
                "Central purchasing block",
            ],
        )

        for row in rows:
            if _text(row["Supplier"]) == requested_id:
                return {
                    "supplier_id": _text(row["Supplier"]),
                    "name": _text(row["Name"]),
                    "name_2": _text(row["Name 2"]),
                    "country": _text(row["Country/Region Key"]),
                    "city": _text(row["City"]),
                    "deletion_flag": _text(
                        row["Central Deletion Flag"]
                    ),
                    "posting_block": _text(
                        row["Central posting block"]
                    ),
                    "purchasing_block": _text(
                        row["Central purchasing block"]
                    ),
                    "source": "LFA1_EXCEL_EXPORT",
                }

        return None
    
    
    def get_purchase_order_context(
        self,
        po_number: str,
        supplier_id: str | None = None,
    ) -> dict[str, Any] | None:
        requested_po = _text(po_number)
        if not requested_po:
            return None

        headers = self._rows(
            "ekko",
            [
                "Purchasing Document",
                "Company Code",
                "Purchasing Doc. Type",
                "Deletion Indicator",
                "Status",
                "Supplier",
                "Currency",
                "Document Date",
            ],
        )

        header = next(
            (
                row
                for row in headers
                if _text(row["Purchasing Document"]) == requested_po
            ),
            None,
        )

        if header is None:
            return None

        resolved_supplier_id = _text(header["Supplier"])
        supplier = (
            self.get_supplier(resolved_supplier_id)
            if resolved_supplier_id
            else None
        )

        item_rows = self._rows(
            "ekpo",
            [
                "Purchasing Document",
                "Item",
                "Deletion Indicator",
                "Short Text",
                "Material",
                "Company Code",
                "Plant",
                "Order Quantity",
                "Net Order Price",
                "Price unit",
                "Net Order Value",
                "Goods Receipt",
                "Invoice Receipt",
                "Delivery Completed",
                "G/L Account",
                "Cost Center",
            ],
        )

        po_lines: list[dict[str, Any]] = []
        po_total = Decimal("0")
        total_is_complete = True

        for row in item_rows:
            if _text(row["Purchasing Document"]) != requested_po:
                continue

            deleted = _text(row["Deletion Indicator"])
            is_active = not bool(deleted)
            line_amount = _decimal(row["Net Order Value"])

            if is_active:
                if line_amount is None:
                    total_is_complete = False
                else:
                    po_total += line_amount

            po_lines.append({
                "item": _text(row["Item"]),
                "description": _text(row["Short Text"]),
                "material": _text(row["Material"]),
                "company_code": _text(row["Company Code"]),
                "plant": _text(row["Plant"]),
                "order_quantity": _text(row["Order Quantity"]),
                "net_order_price": _text(row["Net Order Price"]),
                "price_unit": _text(row["Price unit"]),
                "net_order_value": _text(row["Net Order Value"]),
                "deletion_indicator": deleted,
                "is_active": is_active,
                "goods_receipt_expected": _text(
                    row["Goods Receipt"]
                ),
                "invoice_receipt_expected": _text(
                    row["Invoice Receipt"]
                ),
                "delivery_completed": _text(
                    row["Delivery Completed"]
                ),
                "gl_account": _text(row["G/L Account"]),
                "cost_center": _text(row["Cost Center"]),
            })

        if not po_lines:
            total_is_complete = False

        # Do not confuse the PO's original value with its remaining balance.
        # The balance requires verified GR and invoice-consumption logic.
        return {
            "po_number": requested_po,
            "supplier_id": resolved_supplier_id,
            "supplier": supplier,
            "company_code": _text(header["Company Code"]),
            "document_type": _text(header["Purchasing Doc. Type"]),
            "deletion_indicator": _text(
                header["Deletion Indicator"]
            ),
            "status": _text(header["Status"]),
            "currency": _text(header["Currency"]),
            "document_date": _json_value(header["Document Date"]),
            "total_amount": (
                str(po_total) if total_is_complete else None
            ),
            "available_amount": None,
            "total_amount_is_complete": total_is_complete,
            "lines": po_lines,
            "supplier_lookup_found": supplier is not None,
            "source": "LOCAL_SAP_EXCEL_EXPORT",
            "fetched_at": datetime.now(timezone.utc).isoformat(),
            "snapshot_hash": self._snapshot_hash(),
        }
        
    def get_purchase_order_history(
        self,
        po_number: str,
    ) -> list[dict[str, Any]]:
        requested_po = _text(po_number)
        if not requested_po:
            return []

        rows = self._rows(
            "ekbe",
            [
                "Purchasing Document",
                "Item",
                "Trans./event type",
                "Material Document",
                "Material Doc. Year",
                "PO History Category",
                "Movement Type",
                "Posting Date",
                "Quantity",
                "Amount",
                "Currency",
                "Debit/Credit Ind.",
                "Invoice Value",
                "Service",
                "Package Number",
                "Service line",
            ],
        )

        history = []

        for row in rows:
            if _text(row["Purchasing Document"]) != requested_po:
                continue

            history.append({
                "item": _text(row["Item"]),
                "event_type": _text(row["Trans./event type"]),
                "material_document": _text(row["Material Document"]),
                "material_document_year": _text(
                    row["Material Doc. Year"]
                ),
                "history_category": _text(
                    row["PO History Category"]
                ),
                "movement_type": _text(row["Movement Type"]),
                "posting_date": _json_value(row["Posting Date"]),
                "quantity": _text(row["Quantity"]),
                "amount": _text(row["Amount"]),
                "currency": _text(row["Currency"]),
                "debit_credit_indicator": _text(
                    row["Debit/Credit Ind."]
                ),
                "invoice_value": _text(row["Invoice Value"]),
                "service": _text(row["Service"]),
                "package_number": _text(row["Package Number"]),
                "service_line": _text(row["Service line"]),
            })

        return history
