
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_DATA_DIR = (
    Path(__file__).resolve().parents[2] / "data" / "sap_excel"
)

FILES = {
    "banks": "BNKA - Bank master record.xlsx",
    "cost_centers": "CSKS - Cost Center Master Record.xlsx",
    "cost_center_texts": "CSKT - Cost Center Texts.xlsx",
    "po_assignments": "EKKN - PO account assignment.xlsx",
    "service_lines": "ESLL - Lines of Service Package.xlsx",
    "service_entries": "ESSR - Service Entry Sheet Header Data.xlsx",
    "vendor_banks": "LFBK - Vendor Master (Bank Details).xlsx",
    "gl_accounts": "SKA1 - GL Account Master Data (Chart of Accounts).xlsx",
    "gl_texts": "SKAT - GL Account Master Description.xlsx",
}


def _key(value: Any) -> str:
    """Normalize Excel keys without losing string identifiers."""
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    return str(value).strip()


def _text(value: Any) -> str:
    return _key(value)


def _active_rows(path: Path) -> list[dict[str, Any]]:
    """Load a workbook's first worksheet into dictionaries."""
    if not path.is_file():
        raise FileNotFoundError(f"SAP reference file not found: {path}")

    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)

        headers = next(rows, None)
        if not headers:
            return []

        # Keep duplicate Excel column names distinguishable.
        unique_headers = []
        seen: dict[str, int] = {}

        for index, value in enumerate(headers):
            name = _text(value) or f"column_{index}"
            seen[name] = seen.get(name, 0) + 1

            if seen[name] > 1:
                name = f"{name}__{seen[name]}"

            unique_headers.append(name)

        records = []

        for values in rows:
            if not any(value is not None for value in values):
                continue

            records.append({
                unique_headers[index]: value
                for index, value in enumerate(values)
                if index < len(unique_headers)
            })

        return records
    finally:
        workbook.close()


def _value(row: dict[str, Any], column: str) -> Any:
    """Read an exact workbook column."""
    return row.get(column)


def _matches(
    row: dict[str, Any],
    column: str,
    expected: Any,
) -> bool:
    expected_key = _key(expected)

    return bool(expected_key) and (
        _key(_value(row, column)) == expected_key
    )


class ExcelSAPFinanceReader:
    """
    Read-only reader for SAP finance reference Excel exports.

    This reader does not write to SAP or approve/post invoices.
    It reports reference-data evidence; business decisions remain
    with AP Core and its deterministic validation rules.
    """

    def __init__(
        self,
        data_dir: str | Path | None = None,
    ) -> None:
        self.data_dir = Path(data_dir or DEFAULT_DATA_DIR)

        self._data: dict[str, list[dict[str, Any]]] = {}

        for name, filename in FILES.items():
            self._data[name] = _active_rows(
                self.data_dir / filename
            )

    def dataset_counts(self) -> dict[str, int]:
        """Return loaded row counts for diagnostics."""
        return {
            name: len(rows)
            for name, rows in self._data.items()
        }

    def get_vendor_bank_details(
        self,
        supplier_id: str,
    ) -> dict[str, Any]:
        """
        Return supplier bank records joined to bank master data.

        Multiple bank accounts are preserved. A record's presence
        does not by itself mean that the account is verified.
        """
        supplier_id = _key(supplier_id)

        if not supplier_id:
            return {
                "status": "UNKNOWN",
                "reason": "SUPPLIER_ID_MISSING",
                "records": [],
            }

        vendor_records = [
            row for row in self._data["vendor_banks"]
            if _matches(row, "Supplier", supplier_id)
        ]

        joined = []

        for vendor in vendor_records:
            country = _key(
                _value(vendor, "Bank Country/Region")
            )
            bank_key = _key(_value(vendor, "Bank Key"))

            bank_matches = [
                bank for bank in self._data["banks"]
                if _matches(bank, "Bank Country/Region", country)
                and _matches(bank, "Bank Key", bank_key)
            ]

            joined.append({
                "supplier_id": supplier_id,
                "bank_country": country,
                "bank_key": bank_key,
                "bank_account": _value(vendor, "Bank Account"),
                "account_holder": _value(vendor, "Account holder"),
                "account_name": _value(vendor, "Account Name"),
                "valid_from": _value(vendor, "Valid from"),
                "valid_to": _value(vendor, "Eff.to"),
                "bank_records": [
                    {
                        "bank_name": _value(bank, "Bank Name"),
                        "city": _value(bank, "City"),
                        "swift_bic": _value(bank, "SWIFT/BIC"),
                        "deletion_indicator": _value(
                            bank, "Deletion Indicator"
                        ),
                    }
                    for bank in bank_matches
                ],
                "bank_master_match": bool(bank_matches),
            })

        return {
            "status": "FOUND" if joined else "NOT_FOUND",
            "supplier_id": supplier_id,
            "record_count": len(joined),
            "records": joined,
            "source": [
                FILES["vendor_banks"],
                FILES["banks"],
            ],
        }

    def get_bank(
        self,
        country: str,
        bank_key: str,
    ) -> dict[str, Any]:
        country = _key(country)
        bank_key = _key(bank_key)

        matches = [
            row for row in self._data["banks"]
            if _matches(row, "Bank Country/Region", country)
            and _matches(row, "Bank Key", bank_key)
        ]

        return {
            "status": "FOUND" if matches else "NOT_FOUND",
            "records": matches,
            "source": FILES["banks"],
        }

    def get_po_account_assignments(
        self,
        po_number: str,
        item_number: str | None = None,
    ) -> dict[str, Any]:
        po_number = _key(po_number)

        if not po_number:
            return {
                "status": "UNKNOWN",
                "reason": "PO_NUMBER_MISSING",
                "records": [],
            }

        matches = [
            row for row in self._data["po_assignments"]
            if _matches(row, "Purchasing Document", po_number)
            and (
                item_number is None
                or _matches(row, "Item", item_number)
            )
        ]

        records = []

        for row in matches:
            records.append({
                "po_number": po_number,
                "item_number": _value(row, "Item"),
                "account_assignment_sequence": _value(
                    row, "Seq. No. of Account Assgt"
                ),
                "gl_account": _value(row, "G/L Account"),
                "cost_center": _value(row, "Cost Center"),
                "controlling_area": _value(row, "Controlling Area"),
                "profit_center": _value(row, "Profit Center"),
                "quantity": _value(row, "Quantity"),
                "net_order_value": _value(row, "Net Order Value"),
                "deletion_indicator": _value(
                    row, "Deletion Indicator"
                ),
                "tax_code": _value(row, "Tax Code"),
            })

        return {
            "status": "FOUND" if records else "NOT_FOUND",
            "records": records,
            "source": FILES["po_assignments"],
        }

    def get_service_entry_sheets(
        self,
        po_number: str,
        item_number: str | None = None,
    ) -> dict[str, Any]:
        po_number = _key(po_number)

        if not po_number:
            return {
                "status": "UNKNOWN",
                "reason": "PO_NUMBER_MISSING",
                "records": [],
            }

        matches = [
            row for row in self._data["service_entries"]
            if _matches(row, "Purchasing Document", po_number)
            and (
                item_number is None
                or _matches(row, "Item", item_number)
            )
        ]

        records = []

        for row in matches:
            package_number = _key(
                _value(row, "Package number")
            )

            service_lines = [
                line for line in self._data["service_lines"]
                if package_number
                and _matches(line, "Package number", package_number)
            ]

            records.append({
                "entry_sheet": _value(row, "Entry Sheet"),
                "po_number": po_number,
                "item_number": _value(row, "Item"),
                "package_number": package_number,
                "gross_value": _value(row, "Gross Value"),
                "net_value": _value(row, "Net Value"),
                "currency": _value(row, "Currency"),
                "acceptance": _value(row, "Acceptance"),
                "final_entry": _value(row, "Final Entry"),
                "release_indicator": _value(row, "Release indicator"),
                "release_state": _value(row, "Release State"),
                "deletion_indicator": _value(
                    row, "Deletion indicator"
                ),
                "service_lines": [
                    {
                        "line": _value(line, "Line"),
                        "description": _value(line, "Short Text"),
                        "quantity": _value(line, "Quantity"),
                        "unit": _value(line, "Base Unit of Measure"),
                        "gross_price": _value(line, "Gross Price"),
                        "net_value": _value(line, "Net Value"),
                        "tax_code": _value(line, "Tax Code"),
                        "deletion_indicator": _value(
                            line, "Deletion Indicator"
                        ),
                    }
                    for line in service_lines
                ],
            })

        return {
            "status": "FOUND" if records else "NOT_FOUND",
            "records": records,
            "source": [
                FILES["service_entries"],
                FILES["service_lines"],
            ],
        }

    def get_gl_account(
        self,
        chart_of_accounts: str,
        gl_account: str,
        language: str | None = None,
    ) -> dict[str, Any]:
        chart = _key(chart_of_accounts)
        account = _key(gl_account)

        if not chart or not account:
            return {
                "status": "UNKNOWN",
                "reason": "CHART_OR_GL_ACCOUNT_MISSING",
                "records": [],
            }

        matches = [
            row for row in self._data["gl_accounts"]
            if _matches(row, "Chart of Accounts", chart)
            and _matches(row, "G/L Account", account)
        ]

        descriptions = [
            row for row in self._data["gl_texts"]
            if _matches(row, "Chart of Accounts", chart)
            and _matches(row, "G/L Account", account)
            and (
                language is None
                or _matches(row, "Language Key", language)
            )
        ]

        records = []

        for row in matches:
            records.append({
                "chart_of_accounts": chart,
                "gl_account": account,
                "account_type": _value(row, "G/L Account Type"),
                "account_subtype": _value(row, "G/L Account Subtype"),
                "reconciliation_account": _value(
                    row, "Reconciliation Account"
                ),
                "blocked_for_creation": _value(
                    row, "Blocked for Creation"
                ),
                "blocked_for_posting": _value(
                    row, "Blocked for Posting"
                ),
                "mark_for_deletion": _value(row, "Mark for Deletion"),
                "descriptions": [
                    {
                        "language": _value(text, "Language Key"),
                        "short_text": _value(text, "Short Text"),
                        "long_text": _value(text, "G/L Long Text"),
                    }
                    for text in descriptions
                ],
            })

        return {
            "status": "FOUND" if records else "NOT_FOUND",
            "records": records,
            "source": [
                FILES["gl_accounts"],
                FILES["gl_texts"],
            ],
        }

    def get_cost_center(
        self,
        controlling_area: str,
        cost_center: str,
        company_code: str | None = None,
        language: str | None = None,
    ) -> dict[str, Any]:
        controlling_area = _key(controlling_area)
        cost_center = _key(cost_center)

        if not controlling_area or not cost_center:
            return {
                "status": "UNKNOWN",
                "reason": "CONTROLLING_AREA_OR_COST_CENTER_MISSING",
                "records": [],
            }

        matches = [
            row for row in self._data["cost_centers"]
            if _matches(row, "Controlling Area", controlling_area)
            and _matches(row, "Cost Center", cost_center)
            and (
                company_code is None
                or _matches(row, "Company Code", company_code)
            )
        ]

        descriptions = [
            row for row in self._data["cost_center_texts"]
            if _matches(row, "Controlling Area", controlling_area)
            and _matches(row, "Cost Center", cost_center)
            and (
                language is None
                or _matches(row, "Language Key", language)
            )
        ]

        records = []

        for row in matches:
            records.append({
                "controlling_area": controlling_area,
                "cost_center": cost_center,
                "company_code": _value(row, "Company Code"),
                "valid_from": _value(row, "Valid From"),
                "valid_to": _value(row, "Valid To"),
                "currency": _value(row, "Currency"),
                "profit_center": _value(row, "Profit Center"),
                "department": _value(row, "Department"),
                "descriptions": [
                    {
                        "language": _value(text, "Language Key"),
                        "name": _value(text, "Name"),
                        "description": _value(text, "Description"),
                    }
                    for text in descriptions
                ],
            })

        return {
            "status": "FOUND" if records else "NOT_FOUND",
            "records": records,
            "source": [
                FILES["cost_centers"],
                FILES["cost_center_texts"],
            ],
        }
